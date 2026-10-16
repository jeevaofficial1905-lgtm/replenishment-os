"""
server.py — Fullstack Production Server for Render.com
Serves the ReplenishmentOS dashboard and provides high-speed Excel conversion endpoints.
Compatible with Python 3.8 - 3.14+ (No deprecated modules like cgi).
"""

import os
import sys
import json
import io
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from openpyxl import load_workbook

PORT = int(os.environ.get("PORT", 10000))
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

def extract_file_bytes(headers, body):
    content_type = headers.get('content-type', '')
    if 'multipart/form-data' in content_type and 'boundary=' in content_type:
        boundary = content_type.split('boundary=')[1].strip()
        parts = body.split(f"--{boundary}".encode('latin1'))
        for part in parts:
            if b'filename=' in part and b'\r\n\r\n' in part:
                header_part, file_data = part.split(b'\r\n\r\n', 1)
                if file_data.endswith(b'\r\n'):
                    file_data = file_data[:-2]
                return file_data
    return body

def clean_excel_stream(file_stream):
    """
    Transforms raw 3-tier ERP Excel dump into normalized records.
    """
    wb = load_workbook(filename=file_stream, read_only=True, data_only=True)
    sheet = wb[wb.sheetnames[0]]
    rows_iter = sheet.iter_rows(values_only=True)
    
    r1 = next(rows_iter)
    r2 = next(rows_iter)
    r3 = next(rows_iter)
    
    clean_headers = []
    ebo_idx = 0
    d2c_idx = 0
    lfr_idx = 0
    
    for i, h in enumerate(r3):
        h_str = str(h).strip() if h is not None else ''
        if i == 0: clean_headers.append('Style ID')
        elif i == 1: clean_headers.append('Option ID')
        elif i == 2: clean_headers.append('Description')
        elif i == 3: clean_headers.append('SKU')
        elif i == 4: clean_headers.append('Size')
        elif 5 <= i <= 22:
            mo = h_str.replace('EBO', '').strip()
            yr = '25' if ebo_idx < 9 else '26'
            clean_headers.append(f"EBO {mo}-{yr}")
            ebo_idx += 1
        elif 23 <= i <= 40:
            mo = h_str.replace('D2C', '').strip()
            yr = '25' if d2c_idx < 9 else '26'
            clean_headers.append(f"D2C {mo}-{yr}")
            d2c_idx += 1
        elif 41 <= i <= 57:
            mo = h_str.replace('LFR', '').strip()
            if mo.upper() == 'JULY': mo = 'JUL'
            yr = '25' if lfr_idx < 9 else '26'
            clean_headers.append(f"LFR {mo}-{yr}")
            lfr_idx += 1
        elif h_str == 'WareIQ Gur':
            clean_headers.append('WareIQ Ggn')
        elif 'Pending GRN' in h_str:
            clean_headers.append('Pending GRN Units')
        else:
            clean_headers.append(h_str)
            
    aug_missing = 'AUG' not in clean_headers and 'SEP' in clean_headers
    if aug_missing:
        sep_idx = clean_headers.index('SEP')
        clean_headers.insert(sep_idx, 'AUG')
        
    records = []
    for row in rows_iter:
        if not row or row[0] is None or str(row[0]).strip() == '':
            continue
        row_dict = {}
        col_offset = 0
        for col_idx, col_name in enumerate(clean_headers):
            if aug_missing and col_name == 'AUG':
                row_dict['AUG'] = 0
                col_offset = 1
            else:
                raw_idx = col_idx - col_offset
                val = row[raw_idx] if raw_idx < len(row) else 0
                row_dict[col_name] = 0 if val is None else val
        records.append(row_dict)
        
    return records

class RenderHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=BASE_DIR, **kwargs)

    def do_GET(self):
        # Health check endpoint for Render
        if self.path in ('/health', '/api/health'):
            self.send_response(200)
            self.send_header('Content-Type', 'application/json')
            self.end_headers()
            self.wfile.write(b'{"status":"ok","app":"ReplenishmentOS"}')
            return
        # Root path serves index.html
        if self.path == '/':
            self.path = '/index.html'
        return super().do_GET()

    def do_POST(self):
        # Server-side API endpoint for Excel conversion
        if self.path == '/api/convert':
            try:
                content_length = int(self.headers.get('content-length', 0))
                body = self.rfile.read(content_length)

                if not body:
                    self.send_error(400, "No file uploaded")
                    return

                file_bytes = extract_file_bytes(self.headers, body)
                records = clean_excel_stream(io.BytesIO(file_bytes))
                payload = json.dumps(records).encode('utf-8')

                self.send_response(200)
                self.send_header('Content-Type', 'application/json')
                self.send_header('Content-Length', str(len(payload)))
                self.send_header('Access-Control-Allow-Origin', '*')
                self.end_headers()
                self.wfile.write(payload)
                return
            except Exception as e:
                err_msg = json.dumps({"error": str(e)}).encode('utf-8')
                self.send_response(500)
                self.send_header('Content-Type', 'application/json')
                self.end_headers()
                self.wfile.write(err_msg)
                return

        return super().do_POST()

def run():
    server = ThreadingHTTPServer(('0.0.0.0', PORT), RenderHandler)
    print(f"==================================================")
    print(f"ReplenishmentOS Server running on port {PORT}")
    print(f"Access locally at: http://localhost:{PORT}")
    print(f"Ready for Render.com deployment")
    print(f"==================================================")
    server.serve_forever()

if __name__ == '__main__':
    run()
