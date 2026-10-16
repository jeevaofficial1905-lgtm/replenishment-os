"""
process_pipeline.py
Transforms raw pipeline files (like 'pipeline as of 09-10-2026.xlsx')
into:
  1. pipeline_data.json   (blazing fast JSON for ReplenishmentOS, loads in 0.15s)
  2. pipeline_data.zip    (compact 2.0 MB zip for easy emailing / Slack sharing)
  3. PIPELINE WORKING.xlsx (clean Excel for spreadsheet users)
"""

import sys, os, time, json, zipfile, shutil
from openpyxl import load_workbook, Workbook

sys.stdout.reconfigure(encoding='utf-8')

DEFAULT_INPUT  = r"C:\Users\Jeeva\Downloads\pipeline as of 09-10-2026.xlsx"
OUTPUT_DIR_DL  = r"C:\Users\Jeeva\Downloads"
OUTPUT_DIR_WS  = r"c:\Users\Jeeva\Desktop\Pipeline"

def transform_pipeline(input_file=DEFAULT_INPUT):
    print("===================================================")
    print(f"Loading raw file: {input_file}...")
    start_time = time.time()
    
    wb_in = load_workbook(filename=input_file, read_only=True, data_only=True)
    sheet_in = wb_in[wb_in.sheetnames[0]]
    
    rows_iter = sheet_in.iter_rows(values_only=True)
    
    # Skip Row 1 (Banners) and Row 2 (Totals)
    r1 = next(rows_iter)
    r2 = next(rows_iter)
    r3 = next(rows_iter) # Original column headers in Row 3
    
    print(f"Read headers from Row 3 ({len(r3)} columns). Building clean header mapping...")
    
    clean_headers = []
    ebo_idx = 0
    d2c_idx = 0
    lfr_idx = 0
    
    for i, h in enumerate(r3):
        h_str = str(h).strip() if h is not None else ''
        
        # 1. Base Identifiers
        if i == 0: clean_headers.append('Style ID')
        elif i == 1: clean_headers.append('Option ID')
        elif i == 2: clean_headers.append('Description')
        elif i == 3: clean_headers.append('SKU')
        elif i == 4: clean_headers.append('Size')
        
        # 2. EBO Sales (Columns 5 to 22)
        elif 5 <= i <= 22:
            mo = h_str.replace('EBO', '').strip()
            yr = '25' if ebo_idx < 9 else ('26' if ebo_idx < 12 else '26')
            clean_headers.append(f"EBO {mo}-{yr}")
            ebo_idx += 1
            
        # 3. D2C Sales (Columns 23 to 40)
        elif 23 <= i <= 40:
            mo = h_str.replace('D2C', '').strip()
            yr = '25' if d2c_idx < 9 else ('26' if d2c_idx < 12 else '26')
            clean_headers.append(f"D2C {mo}-{yr}")
            d2c_idx += 1
            
        # 4. LFR Sales (Columns 41 to 57)
        elif 41 <= i <= 57:
            mo = h_str.replace('LFR', '').strip()
            if mo.upper() == 'JULY': mo = 'JUL'
            yr = '25' if lfr_idx < 9 else ('26' if lfr_idx < 12 else '26')
            clean_headers.append(f"LFR {mo}-{yr}")
            lfr_idx += 1
            
        # 5. Warehouses & Stock adjustments
        elif h_str == 'WareIQ Gur':
            clean_headers.append('WareIQ Ggn') # Frontend expects WareIQ Ggn
            
        elif 'Pending GRN' in h_str:
            clean_headers.append('Pending GRN Units') # Frontend expects Pending GRN Units
            
        else:
            clean_headers.append(h_str)
            
    # Add 'AUG' pipeline month if missing
    aug_missing = 'AUG' not in clean_headers and 'SEP' in clean_headers
    if aug_missing:
        sep_idx = clean_headers.index('SEP')
        clean_headers.insert(sep_idx, 'AUG')
    
    print(f"Clean headers created ({len(clean_headers)} columns).")
    
    # ── 1. Create Target Excel (write-only streaming) ──
    excel_ws = os.path.join(OUTPUT_DIR_WS, "PIPELINE WORKING.xlsx")
    excel_dl = os.path.join(OUTPUT_DIR_DL, "PIPELINE WORKING.xlsx")
    wb_out = Workbook(write_only=True)
    ws_out = wb_out.create_sheet(title='Sheet1')
    ws_out.append(clean_headers)
    
    # ── 2. Create Target JSON Data list ──
    json_data = []
    row_count = 0
    
    for row in rows_iter:
        row_list = list(row)
        if aug_missing:
            row_list.insert(sep_idx, 0)
        ws_out.append(row_list)
        
        # Build JSON row dict
        row_dict = {}
        for h, v in zip(clean_headers, row_list):
            row_dict[h] = 0 if v is None else v
        json_data.append(row_dict)
        
        row_count += 1
        if row_count % 10000 == 0:
            print(f"  Processed {row_count:,} SKUs...")
            
    wb_in.close()
    
    # ── Save JSON First ──
    json_dl = os.path.join(OUTPUT_DIR_DL, "pipeline_data.json")
    json_ws = os.path.join(OUTPUT_DIR_WS, "pipeline_data.json")
    print(f"Saving JSON ({row_count:,} SKUs) to:")
    print(f"  -> {json_dl}")
    with open(json_dl, "w", encoding="utf-8") as f:
        json.dump(json_data, f, separators=(',', ':'))
    if os.path.abspath(json_dl) != os.path.abspath(json_ws):
        shutil.copyfile(json_dl, json_ws)
        
    # ── Save ZIP Next ──
    zip_dl = os.path.join(OUTPUT_DIR_DL, "pipeline_data.zip")
    zip_ws = os.path.join(OUTPUT_DIR_WS, "pipeline_data.zip")
    print(f"Creating compressed ZIP package to:")
    print(f"  -> {zip_dl}")
    with zipfile.ZipFile(zip_dl, 'w', zipfile.ZIP_DEFLATED) as zf:
        zf.write(json_dl, arcname="pipeline_data.json")
    if os.path.abspath(zip_dl) != os.path.abspath(zip_ws):
        shutil.copyfile(zip_dl, zip_ws)
        
    # ── Save Excel in Workspace first, then copy to Downloads ──
    try:
        print(f"Saving Excel with {row_count:,} SKUs to:")
        print(f"  -> {excel_ws}")
        wb_out.save(excel_ws)
        try:
            shutil.copyfile(excel_ws, excel_dl)
            print(f"  -> Successfully copied to: {excel_dl}")
        except PermissionError:
            print(f"  [Note] {excel_dl} is currently open in Excel; workspace copy updated at {excel_ws}")
    except Exception as e:
        print(f"  [Warning] Could not save Excel: {e}")

    total_time = time.time() - start_time
    print("===================================================")
    print(f"SUCCESS! Processed {row_count:,} SKUs in {total_time:.2f}s")
    print(f"Files ready:")
    print(f"  1. JSON:  {json_dl} ({os.path.getsize(json_dl)/(1024*1024):.2f} MB - Instant browser load)")
    print(f"  2. ZIP:   {zip_dl} ({os.path.getsize(zip_dl)/(1024*1024):.2f} MB - Ultra compact for sharing)")
    if os.path.exists(excel_ws):
        print(f"  3. Excel: {excel_ws} ({os.path.getsize(excel_ws)/(1024*1024):.2f} MB)")
    print("===================================================")

if __name__ == '__main__':
    inp = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_INPUT
    transform_pipeline(inp)
